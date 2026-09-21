from __future__ import annotations

import asyncio
import os
import re
import tempfile
import uuid
from collections.abc import Callable
from contextlib import suppress
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from extrais_leads.models import Search, SearchResult
from extrais_leads.models.base import utc_now
from extrais_leads.repositories.searches import SearchRepository

EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_CELL_CHARACTERS = 32_767

EXCEL_HEADERS = (
    "Posição",
    "Nome da empresa",
    "Telefone",
    "WhatsApp",
    "Status do WhatsApp",
    "Evidência do WhatsApp",
    "URL da evidência do WhatsApp",
    "Fontes",
    "URLs das fontes",
    "Endereço",
    "Cidade",
    "Estado/Região",
    "Categoria/Segmento",
    "CNPJ",
    "Website",
    "Instagram",
    "Confiança dos dados (%)",
    "Correspondência da categoria",
    "Confiança da qualificação (%)",
    "Método da qualificação",
    "Motivo da qualificação",
    "Data da coleta (UTC)",
    "ID da empresa",
    "ID do resultado",
    "ID da pesquisa",
)

_COLUMN_WIDTHS = (
    10,
    34,
    18,
    18,
    22,
    55,
    45,
    32,
    45,
    42,
    22,
    18,
    30,
    20,
    38,
    34,
    24,
    28,
    30,
    28,
    55,
    24,
    38,
    38,
    38,
)
_ILLEGAL_XML_CHARACTERS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_SENSITIVE_QUERY_TERMS = frozenset(
    {
        "api",
        "apikey",
        "access",
        "bearer",
        "auth",
        "authorization",
        "credential",
        "credentials",
        "key",
        "password",
        "refresh",
        "secret",
        "signature",
        "token",
    }
)
_SENSITIVE_QUERY_COMPACT_KEYS = frozenset(
    {
        "accesstoken",
        "apikey",
        "refreshtoken",
    }
)


class ExcelExportError(RuntimeError):
    """A sanitized export failure safe to map to a generic HTTP response."""


@dataclass(frozen=True, slots=True)
class ExcelExportArtifact:
    path: Path
    filename: str
    row_count: int

    def cleanup(self) -> None:
        with suppress(FileNotFoundError):
            self.path.unlink()


@dataclass(slots=True)
class _WorksheetState:
    sheet: Any
    data_rows: int = 0


class ExcelExportService:
    """Generate bounded-memory Excel exports from persisted search results."""

    def __init__(
        self,
        repository: SearchRepository | None = None,
        *,
        batch_size: int = 500,
        temp_directory: Path | None = None,
        clock: Callable[[], datetime] = utc_now,
    ) -> None:
        if batch_size < 1:
            raise ValueError("batch_size must be positive")
        self._repository = repository or SearchRepository()
        self._batch_size = batch_size
        self._temp_directory = temp_directory
        self._clock = clock

    async def export(
        self,
        session: AsyncSession,
        search: Search,
    ) -> ExcelExportArtifact:
        """Write all current persisted results to a temporary XLSX artifact."""

        path = self._new_temp_path()
        workbook: Workbook | None = None
        try:
            workbook = Workbook(write_only=True)
            workbook.properties.creator = "Extrai Leads"
            workbook.properties.title = "Empresas encontradas"
            workbook.properties.subject = "Exportação auditável de resultados de pesquisa"
            state = self._new_worksheet(workbook, 1)

            offset = 0
            row_count = 0
            while True:
                batch = await self._repository.list_result_batch(
                    session,
                    search.id,
                    offset=offset,
                    limit=self._batch_size,
                )
                if not batch:
                    break

                for result in batch:
                    if state.data_rows >= EXCEL_MAX_ROWS - 1:
                        self._finalize_worksheet(state)
                        state = self._new_worksheet(workbook, len(workbook.worksheets) + 1)
                    state.sheet.append(_result_row(search.id, result))
                    state.data_rows += 1
                    row_count += 1

                offset += len(batch)
                if len(batch) < self._batch_size:
                    break

            self._finalize_worksheet(state)
            await self._save_workbook(workbook, path)
            if not path.is_file() or path.stat().st_size == 0:
                raise ExcelExportError("excel writer produced no file")
            return ExcelExportArtifact(
                path=path,
                filename=_download_filename(search.id, self._clock()),
                row_count=row_count,
            )
        except asyncio.CancelledError:
            _remove_file(path)
            raise
        except ExcelExportError:
            _remove_file(path)
            raise
        except Exception as exc:
            _remove_file(path)
            raise ExcelExportError("excel export generation failed") from exc
        finally:
            if workbook is not None:
                with suppress(Exception):
                    workbook.close()

    async def _save_workbook(self, workbook: Workbook, path: Path) -> None:
        await asyncio.to_thread(workbook.save, path)

    def _new_temp_path(self) -> Path:
        file_descriptor, raw_path = tempfile.mkstemp(
            prefix="extrais-leads-",
            suffix=".xlsx",
            dir=self._temp_directory,
        )
        os.close(file_descriptor)
        return Path(raw_path)

    @staticmethod
    def _new_worksheet(workbook: Workbook, number: int) -> _WorksheetState:
        title = "Empresas" if number == 1 else f"Empresas {number}"
        sheet = workbook.create_sheet(title=title)
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        for column, width in enumerate(_COLUMN_WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.append(_header_cells(sheet))
        return _WorksheetState(sheet=sheet)

    @staticmethod
    def _finalize_worksheet(state: _WorksheetState) -> None:
        last_column = get_column_letter(len(EXCEL_HEADERS))
        state.sheet.auto_filter.ref = f"A1:{last_column}{state.data_rows + 1}"


def _header_cells(sheet: Any) -> list[WriteOnlyCell]:
    cells: list[WriteOnlyCell] = []
    fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for title in EXCEL_HEADERS:
        cell = WriteOnlyCell(sheet, value=title)
        cell.fill = fill
        cell.font = font
        cell.alignment = alignment
        cells.append(cell)
    return cells


def _result_row(search_id: uuid.UUID, result: SearchResult) -> tuple[object, ...]:
    company = result.company
    source_labels = sorted(
        {f"{item.source.display_name} [{item.source.provider_key}]" for item in result.sources},
        key=str.casefold,
    )
    source_urls = sorted(
        {
            safe_url
            for item in result.sources
            if (safe_url := _safe_export_url(item.source_url)) is not None
        }
    )
    whatsapp_evidence = sorted(
        (item for item in result.contact_evidences if item.contact_type == "whatsapp"),
        key=lambda item: (
            not item.official_source,
            item.evidence_type,
            item.source_provider,
            item.source_url,
            item.normalized_value,
        ),
    )
    evidence_descriptions = [_whatsapp_evidence_text(item) for item in whatsapp_evidence]
    evidence_urls = sorted(
        {
            safe_url
            for item in whatsapp_evidence
            if (safe_url := _safe_export_url(item.source_url)) is not None
        }
    )

    values: tuple[object, ...] = (
        result.rank,
        company.name,
        company.phone,
        company.whatsapp,
        company.whatsapp_status.value,
        _join_lines(evidence_descriptions),
        _join_lines(evidence_urls),
        _join_lines(source_labels),
        _join_lines(source_urls),
        company.address,
        company.city,
        company.state,
        company.category,
        company.cnpj,
        _safe_export_url(company.website),
        _safe_export_url(company.instagram),
        result.confidence,
        _optional_boolean(result.category_match),
        result.qualification_confidence,
        result.qualification_method,
        result.qualification_reason,
        _utc_text(result.collected_at),
        str(company.id),
        str(result.id),
        str(search_id),
    )
    return tuple(_safe_cell(value) for value in values)


def _whatsapp_evidence_text(evidence: Any) -> str:
    origin = "oficial" if evidence.official_source else "terceiro"
    parts = [
        evidence.evidence_type,
        evidence.source_provider,
        origin,
        evidence.normalized_value,
    ]
    if evidence.excerpt:
        parts.append(evidence.excerpt)
    return " | ".join(parts)


def _optional_boolean(value: bool | None) -> str | None:
    if value is None:
        return None
    return "sim" if value else "não"


def _join_lines(values: list[str]) -> str | None:
    return "\n".join(values) if values else None


def _utc_text(value: datetime | None) -> str | None:
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(UTC).isoformat(timespec="seconds").replace("+00:00", "Z")


def _safe_cell(value: object) -> object:
    if value is None or isinstance(value, int | float | bool):
        return value
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ")
    text = _ILLEGAL_XML_CHARACTERS.sub("", text)
    if text.lstrip().startswith(_FORMULA_PREFIXES):
        text = "'" + text
    return text[:EXCEL_MAX_CELL_CHARACTERS]


def _safe_export_url(value: str | None) -> str | None:
    if not value:
        return None
    try:
        parsed = urlsplit(value)
        port = parsed.port
    except ValueError:
        return None
    if parsed.scheme.casefold() not in {"http", "https"} or not parsed.hostname:
        return None

    hostname = parsed.hostname.casefold().rstrip(".")
    netloc = f"[{hostname}]" if ":" in hostname else hostname
    if port is not None:
        netloc = f"{netloc}:{port}"
    query = urlencode(
        [
            (key, item)
            for key, item in parse_qsl(parsed.query, keep_blank_values=True)
            if not _is_sensitive_query_key(key)
        ],
        doseq=True,
    )
    return urlunsplit((parsed.scheme.casefold(), netloc, parsed.path, query, ""))


def _is_sensitive_query_key(value: str) -> bool:
    normalized = value.casefold()
    if re.sub(r"[^a-z0-9]+", "", normalized) in _SENSITIVE_QUERY_COMPACT_KEYS:
        return True
    return bool(set(re.findall(r"[a-z0-9]+", normalized)) & _SENSITIVE_QUERY_TERMS)


def _download_filename(search_id: uuid.UUID, generated_at: datetime) -> str:
    if generated_at.tzinfo is None:
        generated_at = generated_at.replace(tzinfo=UTC)
    timestamp = generated_at.astimezone(UTC).strftime("%Y%m%dT%H%M%SZ")
    return f"extrais_leads_{search_id}_{timestamp}.xlsx"


def _remove_file(path: Path) -> None:
    with suppress(FileNotFoundError):
        path.unlink()
