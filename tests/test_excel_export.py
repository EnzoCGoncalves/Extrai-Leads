from __future__ import annotations

import asyncio
import uuid
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

import pytest
from fastapi import FastAPI
from httpx import AsyncClient
from openpyxl import load_workbook

from extrais_leads.models import (
    Company,
    ContactEvidence,
    ResultSource,
    Search,
    SearchResult,
    Source,
)
from extrais_leads.models.enums import SearchStage, SearchStatus, SourceType, WhatsAppStatus
from extrais_leads.services.excel_export import (
    EXCEL_CONTENT_TYPE,
    EXCEL_HEADERS,
    ExcelExportService,
)


async def _finished_search(
    client: AsyncClient,
    app: FastAPI,
    *,
    count: int,
    whatsapp_statuses: list[WhatsAppStatus] | None = None,
    first_name: str = "Clínica São José & Filhos",
    first_address: str = "Rua das Flores, 123",
    sensitive_urls: bool = False,
) -> uuid.UUID:
    created = await client.post(
        "/api/v1/searches",
        json={"query": "Clínicas odontológicas em Campinas"},
    )
    search_id = uuid.UUID(created.json()["id"])

    async with app.state.database.session() as session:
        search = await session.get(Search, search_id)
        if search is None:  # pragma: no cover - defensive assertion with a clearer failure
            raise AssertionError("search was not persisted")

        search.status = SearchStatus.COMPLETED
        search.stage = SearchStage.COMPLETED
        search.progress_percent = 100
        search.results_count = count
        search.whatsapp_count = sum(
            status is not WhatsAppStatus.NOT_FOUND for status in whatsapp_statuses or []
        )
        search.confirmed_whatsapp_count = sum(
            status is WhatsAppStatus.CONFIRMED for status in whatsapp_statuses or []
        )
        search.completed_at = datetime(2026, 9, 21, 15, 30, tzinfo=UTC)

        source = Source(
            provider_key="public_directory",
            display_name="Diretório Público",
            source_type=SourceType.DIRECTORY,
        )
        session.add(source)
        for index in range(count):
            whatsapp_status = (
                whatsapp_statuses[index]
                if whatsapp_statuses and index < len(whatsapp_statuses)
                else WhatsAppStatus.NOT_FOUND
            )
            whatsapp = (
                f"551199999{index:04d}" if whatsapp_status is not WhatsAppStatus.NOT_FOUND else None
            )
            company = Company(
                name=first_name if index == 0 else f"Empresa {index + 1:04d}",
                normalized_name=f"empresa {index + 1}",
                name_fingerprint=f"{index + 1:064x}",
                phone=f"55113333{index:04d}",
                whatsapp=whatsapp,
                whatsapp_status=whatsapp_status,
                address=first_address if index == 0 else f"Rua {index + 1}, 100",
                city="Campinas",
                state="SP",
                category="Clínica odontológica",
                website=(
                    "https://user:password@empresa.example/?api_key=PRIVATE&ref=export"
                    if sensitive_urls and index == 0
                    else f"https://empresa-{index + 1}.example/"
                ),
                instagram="https://instagram.com/clinica.sao.jose" if index == 0 else None,
                cnpj="04252011000110" if index == 0 else None,
                confidence=93,
            )
            result = SearchResult(
                search_id=search_id,
                company=company,
                rank=index + 1,
                confidence=93,
                category_match=True,
                qualification_confidence=92,
                qualification_method="deterministic",
                qualification_reason="A página pública indica atuação em odontologia.",
                collected_at=datetime(2026, 9, 20, 18, 15, tzinfo=UTC),
            )
            source_url = (
                "https://user:password@source.example/company?token=SECRET&monkey=banana&ref=public"
                if sensitive_urls and index == 0
                else f"https://source.example/company/{index + 1}"
            )
            session.add(
                ResultSource(
                    search_result=result,
                    source=source,
                    source_url=source_url,
                    evidence={"api_key": "MUST_NOT_BE_EXPORTED"},
                )
            )
            if whatsapp is not None:
                session.add(
                    ContactEvidence(
                        search_result=result,
                        contact_type="whatsapp",
                        normalized_value=whatsapp,
                        evidence_type=(
                            "direct_link"
                            if whatsapp_status is WhatsAppStatus.CONFIRMED
                            else "structured_field"
                        ),
                        source_provider=(
                            "official_website"
                            if whatsapp_status is WhatsAppStatus.CONFIRMED
                            else "openstreetmap"
                        ),
                        source_url=source_url,
                        official_source=whatsapp_status is WhatsAppStatus.CONFIRMED,
                        excerpt="Fale conosco pelo WhatsApp",
                        details={"secret": "MUST_NOT_BE_EXPORTED"},
                    )
                )
        await session.commit()
    return search_id


def _sheet_rows(content: bytes) -> list[tuple[object, ...]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        return list(workbook["Empresas"].iter_rows(values_only=True))
    finally:
        workbook.close()


@pytest.mark.asyncio
async def test_export_endpoint_generates_professional_xlsx_from_database(
    client: AsyncClient,
    test_app: FastAPI,
    tmp_path: Path,
) -> None:
    test_app.state.excel_export_service = ExcelExportService(
        batch_size=10,
        temp_directory=tmp_path,
        clock=lambda: datetime(2026, 9, 21, 16, 45, tzinfo=UTC),
    )
    search_id = await _finished_search(
        client,
        test_app,
        count=1,
        whatsapp_statuses=[WhatsAppStatus.CONFIRMED],
        sensitive_urls=True,
    )

    response = await client.get(f"/api/v1/searches/{search_id}/export.xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"] == EXCEL_CONTENT_TYPE
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["x-exported-rows"] == "1"
    assert response.headers["content-disposition"] == (
        f'attachment; filename="extrais_leads_{search_id}_20260921T164500Z.xlsx"'
    )
    assert response.content.startswith(b"PK")

    rows = _sheet_rows(response.content)
    assert rows[0] == EXCEL_HEADERS
    exported = dict(zip(EXCEL_HEADERS, rows[1], strict=True))
    assert exported["Nome da empresa"] == "Clínica São José & Filhos"
    assert exported["Telefone"] == "551133330000"
    assert exported["WhatsApp"] == "5511999990000"
    assert exported["Status do WhatsApp"] == "confirmed"
    assert "direct_link | official_website | oficial" in exported["Evidência do WhatsApp"]
    assert exported["Fontes"] == "Diretório Público [public_directory]"
    assert exported["URLs das fontes"] == (
        "https://source.example/company?monkey=banana&ref=public"
    )
    assert exported["Website"] == "https://empresa.example/?ref=export"
    assert exported["Confiança dos dados (%)"] == 93
    assert exported["Correspondência da categoria"] == "sim"
    assert exported["Data da coleta (UTC)"] == "2026-09-20T18:15:00Z"
    assert exported["ID da pesquisa"] == str(search_id)
    assert "PRIVATE" not in str(rows)
    assert "SECRET" not in str(rows)
    assert "MUST_NOT_BE_EXPORTED" not in str(rows)
    assert await asyncio.to_thread(lambda: list(tmp_path.glob("*.xlsx"))) == []


@pytest.mark.asyncio
async def test_export_preserves_all_whatsapp_statuses_without_inference(
    client: AsyncClient,
    test_app: FastAPI,
) -> None:
    search_id = await _finished_search(
        client,
        test_app,
        count=3,
        whatsapp_statuses=[
            WhatsAppStatus.CONFIRMED,
            WhatsAppStatus.UNCONFIRMED,
            WhatsAppStatus.NOT_FOUND,
        ],
    )

    response = await client.get(f"/api/v1/searches/{search_id}/export.xlsx")
    rows = _sheet_rows(response.content)
    status_column = EXCEL_HEADERS.index("Status do WhatsApp")
    whatsapp_column = EXCEL_HEADERS.index("WhatsApp")

    assert [row[status_column] for row in rows[1:]] == [
        "confirmed",
        "unconfirmed",
        "not_found",
    ]
    assert rows[3][whatsapp_column] is None
    assert rows[3][EXCEL_HEADERS.index("Telefone")] == "551133330002"


@pytest.mark.asyncio
async def test_export_empty_completed_search_contains_headers_only(
    client: AsyncClient,
    test_app: FastAPI,
) -> None:
    search_id = await _finished_search(client, test_app, count=0)

    response = await client.get(f"/api/v1/searches/{search_id}/export.xlsx")

    assert response.status_code == 200
    assert response.headers["x-exported-rows"] == "0"
    assert _sheet_rows(response.content) == [EXCEL_HEADERS]


@pytest.mark.asyncio
async def test_export_reads_more_than_api_page_limit_in_bounded_batches(
    client: AsyncClient,
    test_app: FastAPI,
) -> None:
    test_app.state.excel_export_service = ExcelExportService(batch_size=37)
    search_id = await _finished_search(client, test_app, count=513)

    response = await client.get(f"/api/v1/searches/{search_id}/export.xlsx")
    rows = _sheet_rows(response.content)

    assert response.status_code == 200
    assert response.headers["x-exported-rows"] == "513"
    assert len(rows) == 514
    assert rows[-1][EXCEL_HEADERS.index("Nome da empresa")] == "Empresa 0513"


@pytest.mark.asyncio
async def test_export_preserves_unicode_and_neutralizes_excel_formulas(
    client: AsyncClient,
    test_app: FastAPI,
) -> None:
    search_id = await _finished_search(
        client,
        test_app,
        count=1,
        first_name='=HYPERLINK("https://evil.example";"Clínica Ção")',
        first_address="+SUM(1;1) — São Paulo",
    )

    response = await client.get(f"/api/v1/searches/{search_id}/export.xlsx")
    exported = dict(zip(EXCEL_HEADERS, _sheet_rows(response.content)[1], strict=True))

    assert exported["Nome da empresa"].startswith("'=HYPERLINK")
    assert "Clínica Ção" in exported["Nome da empresa"]
    assert exported["Endereço"] == "'+SUM(1;1) — São Paulo"
    assert exported["Categoria/Segmento"] == "Clínica odontológica"


@pytest.mark.asyncio
async def test_export_rejects_non_terminal_search_and_unknown_search(client: AsyncClient) -> None:
    created = await client.post(
        "/api/v1/searches",
        json={"query": "Restaurantes em Campinas"},
    )
    search_id = created.json()["id"]

    processing = await client.get(f"/api/v1/searches/{search_id}/export.xlsx")
    unknown = await client.get("/api/v1/searches/00000000-0000-0000-0000-000000000000/export.xlsx")

    assert processing.status_code == 409
    assert processing.json() == {"detail": "search is still processing"}
    assert unknown.status_code == 404
    assert unknown.json() == {"detail": "search not found"}


class _FailingExcelExportService(ExcelExportService):
    async def _save_workbook(self, workbook: object, path: Path) -> None:
        raise OSError("internal path and SECRET must not reach the API")


@pytest.mark.asyncio
async def test_export_failure_is_sanitized_and_removes_partial_file(
    client: AsyncClient,
    test_app: FastAPI,
    tmp_path: Path,
) -> None:
    test_app.state.excel_export_service = _FailingExcelExportService(temp_directory=tmp_path)
    search_id = await _finished_search(client, test_app, count=0)

    response = await client.get(f"/api/v1/searches/{search_id}/export.xlsx")

    assert response.status_code == 500
    assert response.json() == {"detail": "could not generate Excel export"}
    assert "SECRET" not in response.text
    assert await asyncio.to_thread(lambda: list(tmp_path.glob("*.xlsx"))) == []
